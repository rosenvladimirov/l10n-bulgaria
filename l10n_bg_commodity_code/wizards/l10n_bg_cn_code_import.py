# Copyright 2026 Rosen Vladimirov
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).
"""Импорт на КН за година от CSV или XLSX.

Очакваните източници (виж README):
  * НАП, Приложение 2 към SAF-T документацията — листът NC8_TARIC;
  * НСИ — CN_<година>.xlsx (Комбинирана номенклатура).

Парсването е нарочно търпимо към формата:
  * заглавният ред се търси сред първите 30 реда по ключови думи
    (code/CN8/код/КН …, description/описание/наименование …,
    supplementary unit/допълнителна мярка …);
  * без разпознат заглавен ред кодът е първата клетка, която след махане
    на интервалите и точките е точно 8 цифри, а описанието — най-дългият
    текст в реда;
  * клетка със 7 цифри (число или текст — Excel/CSV е изял водещата нула)
    се допълва с водеща нула до 8 — единственото „допълване“; безопасно е,
    защото в КН няма 7-цифрени нива (глави 2, позиции 4, подпозиции 6);
  * редовете за глави/позиции (2, 4 или 6 цифри) и 10-цифрените TARIC
    подразделения се пропускат и се броят.

openpyxl е по желание: без него се приема само CSV.
"""
import base64
import csv
import io
import re

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:  # openpyxl е по желание — без него остава само CSV
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

_CN8_RE = re.compile(r"^\d{8}$")
_SEPARATORS_RE = re.compile(r"[\s. ]+")

# Ключови думи за колоните (сравняват се с малки букви, без интервали в края)
_HEADER_KEYS = {
    "code": ("cn8", "cn 8", "cn code", "cn_code", "cn-code", "nc8", "kn8",
             "кн8", "кн 8", "код по кн", "код кн", "код", "code", "кн", "cn"),
    "name": ("description", "описание", "наименование", "designation",
             "name", "текст", "стока", "goods"),
    "su": ("supplementary unit", "supplementary_unit", "suppl. unit",
           "suppl unit", "допълнителна мярка", "доп. мярка", "доп мярка",
           "su", "дм", "мярка", "unit"),
}
# Познати допълнителни мерки (за евристиката без заглавен ред)
_KNOWN_SU = {
    "p/st", "100 p/st", "1 000 p/st", "1000 p/st", "pa", "kg", "g", "l",
    "m", "m2", "m3", "1 000 m3", "c/k", "ce/el", "ct/l", "gi f / s",
    "l alc. 100 %", "1 000 kwh", "tj", "t. co2", "kg/net eda",
}
_EMPTY_SU = {"", "-", "—", "–", "n/a", "none"}


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _cell_cn8(value):
    """КН8 от клетка или '' (числата със 7 цифри — изядена водеща нула)."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return ""
        digits = str(int(value))
        if len(digits) == 7:
            digits = digits.zfill(8)
        return digits if _CN8_RE.match(digits) else ""
    cleaned = _SEPARATORS_RE.sub("", str(value))
    # текст със 7 цифри — CSV, изнесен от Excel, който вече е изял водещата нула
    if len(cleaned) == 7 and cleaned.isdigit():
        cleaned = cleaned.zfill(8)
    return cleaned if _CN8_RE.match(cleaned) else ""


def _match_role(text, role):
    text = text.lower().strip().rstrip(":")
    if not text:
        return 0
    keys = _HEADER_KEYS[role]
    if text in keys:
        return 2  # точно съвпадение
    for key in keys:
        # къси ключове („кн“, „su“, „cn“) само като отделна дума
        if len(key) <= 3:
            if re.search(r"(^|[\s_/(-])" + re.escape(key) + r"($|[\s_/)-])", text):
                return 1
        elif key in text:
            return 1
    return 0


class L10nBgCnCodeImport(models.TransientModel):
    _name = "l10n.bg.cn.code.import"
    _description = "Import Combined Nomenclature Codes"

    file = fields.Binary(string="File", required=True)
    filename = fields.Char(string="File Name")
    year = fields.Integer(
        string="Year",
        required=True,
        default=lambda self: fields.Date.context_today(self).year,
        help="Year of the Combined Nomenclature contained in the file.",
    )
    sheet_name = fields.Char(
        string="Sheet",
        help="XLSX only: name of the sheet to read. When empty, a sheet "
             "named like NC8_TARIC or CN_<year> is preferred, otherwise "
             "the first sheet.",
    )
    description_lang = fields.Selection(
        selection="_get_languages",
        string="Description Language",
        default=lambda self: self._default_description_lang(),
        help="Language in which the descriptions of the file are written.",
    )
    archive_missing = fields.Boolean(
        string="Archive Codes Missing from the File",
        help="Archive the codes of this year that are not present in the "
             "file (use when the file is the complete nomenclature).",
    )

    @api.model
    def _get_languages(self):
        return self.env["res.lang"].get_installed()

    @api.model
    def _default_description_lang(self):
        installed = dict(self.env["res.lang"].get_installed())
        if "bg_BG" in installed:
            return "bg_BG"
        return self.env.lang or "en_US"

    # ------------------------------------------------------------------
    # Четене на файла
    # ------------------------------------------------------------------
    def _read_rows(self):
        self.ensure_one()
        content = base64.b64decode(self.file or b"")
        if not content:
            raise UserError(_("The file is empty."))
        name = (self.filename or "").lower()
        is_xlsx = content[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm"))
        if is_xlsx:
            return self._read_xlsx(content)
        return self._read_csv(content)

    def _read_xlsx(self, content):
        if openpyxl is None:
            raise UserError(_(
                "Reading XLSX files requires the Python library openpyxl. "
                "Install it or save the sheet as CSV and import the CSV."
            ))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise UserError(_("The XLSX file cannot be read: %s", exc)) from exc
        sheet = None
        if self.sheet_name:
            if self.sheet_name not in workbook.sheetnames:
                raise UserError(_(
                    "Sheet '%(sheet)s' not found. Available sheets: %(sheets)s",
                    sheet=self.sheet_name,
                    sheets=", ".join(workbook.sheetnames),
                ))
            sheet = workbook[self.sheet_name]
        else:
            for candidate in workbook.sheetnames:
                upper = candidate.upper().replace(" ", "_")
                if upper.startswith(("NC8", "CN_", "CN8", "KN8")):
                    sheet = workbook[candidate]
                    break
            sheet = sheet or workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    def _read_csv(self, content):
        text = None
        for encoding in ("utf-8-sig", "cp1251", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if text.count(";") > text.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]

    # ------------------------------------------------------------------
    # Разпознаване на колоните
    # ------------------------------------------------------------------
    @api.model
    def _detect_header(self, rows):
        """Връща (индекс на заглавния ред, {роля: колона}) или (None, {})."""
        for index, row in enumerate(rows[:30]):
            texts = [_cell_text(c) for c in row]
            if any(_cell_cn8(c) for c in row):
                continue
            columns = {}
            for role in ("code", "su", "name"):
                best, best_score = None, 0
                for col, text in enumerate(texts):
                    if col in columns.values():
                        continue
                    score = _match_role(text, role)
                    if score > best_score:
                        best, best_score = col, score
                if best is not None:
                    columns[role] = best
            if "code" in columns and len(columns) >= 2:
                return index, columns
        return None, {}

    @api.model
    def _parse_rows(self, rows):
        """Връща ({код: (описание, мярка)}, брой пропуснати редове)."""
        header_index, columns = self._detect_header(rows)
        data_rows = rows[header_index + 1:] if header_index is not None else rows
        result, skipped = {}, 0
        for row in data_rows:
            if not row or all(_cell_text(c) == "" for c in row):
                continue
            code, code_col = "", None
            if "code" in columns and columns["code"] < len(row):
                code, code_col = _cell_cn8(row[columns["code"]]), columns["code"]
            if not code:
                for col, cell in enumerate(row):
                    code = _cell_cn8(cell)
                    if code:
                        code_col = col
                        break
            if not code:
                skipped += 1
                continue
            name = ""
            if "name" in columns and columns["name"] < len(row):
                name = _cell_text(row[columns["name"]])
            if not name:
                texts = [
                    _cell_text(c) for col, c in enumerate(row)
                    if col != code_col and not _cell_cn8(c)
                ]
                texts = [t for t in texts if t.lower() not in _KNOWN_SU]
                name = max(texts, key=len) if texts else ""
            su = ""
            if "su" in columns and columns["su"] < len(row):
                su = _cell_text(row[columns["su"]])
            else:
                for cell in row:
                    text = _cell_text(cell)
                    if text.lower() in _KNOWN_SU:
                        su = text
                        break
            if su.lower() in _EMPTY_SU:
                su = ""
            result[code] = (name, su)
        return result, skipped

    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        parsed, skipped = self._parse_rows(self._read_rows())
        if not parsed:
            raise UserError(_(
                "No 8-digit CN code was found in the file. Check the "
                "expected format in the module README."
            ))
        CnCode = self.env["l10n.bg.cn.code"].with_context(
            lang=self.description_lang or self.env.lang, active_test=False,
        )
        existing = {r.code: r for r in CnCode.search([("year", "=", self.year)])}
        to_create, updated = [], 0
        for code, (name, su) in parsed.items():
            record = existing.get(code)
            if not record:
                to_create.append({
                    "code": code, "year": self.year,
                    "name": name or False, "supplementary_unit": su or False,
                })
                continue
            vals = {}
            if name and record.name != name:
                vals["name"] = name
            if (record.supplementary_unit or "") != su:
                vals["supplementary_unit"] = su or False
            if not record.active:
                vals["active"] = True
            if vals:
                record.write(vals)
                updated += 1
        if to_create:
            CnCode.create(to_create)
        archived = 0
        if self.archive_missing:
            missing = CnCode.browse(
                r.id for code, r in existing.items() if code not in parsed and r.active
            )
            missing.write({"active": False})
            archived = len(missing)
        message = _(
            "CN %(year)s: %(created)s created, %(updated)s updated, "
            "%(archived)s archived, %(skipped)s row(s) skipped.",
            year=self.year, created=len(to_create), updated=updated,
            archived=archived, skipped=skipped,
        )
        action = self.env["ir.actions.act_window"]._for_xml_id(
            "l10n_bg_commodity_code.action_l10n_bg_cn_code"
        )
        action["domain"] = [("year", "=", self.year)]
        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Combined Nomenclature imported"),
                "message": message,
                "type": "success",
                "sticky": False,
                "next": action,
            },
        }
